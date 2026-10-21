"""Excel / CSV import for COD Parcels and Courier Remittances.

Flow used by the Desk dialogs: preview_import (nothing saved) -> run_import (parcels are upserted,
remittances are created, submitted and matched). Files above BACKGROUND_ROWS rows run as a
background job. Every import writes one Sync Log row (Excel / Excel Import).
"""

import csv
import datetime
import json
import re

import frappe
from frappe import _
from frappe.utils import cint, flt, get_datetime, getdate, nowdate

from codsync.codsync.engine import normalize_awb

IMPORT_ROLES = ("System Manager", "CODSync Manager", "CODSync Operator")
BACKGROUND_ROWS = 500
PREVIEW_ROWS = 5
MAX_LISTED_ERRORS = 50
MAX_DETAIL_ROWS = 500

PARCEL_COLUMNS = [
	"order_id",
	"awb",
	"courier",
	"customer_name",
	"phone",
	"city",
	"pincode",
	"address",
	"email",
	"cod_amount",
	"order_total",
	"payment_mode",
	"status",
	"delivered_date",
	"order_date",
	"declared_weight",
	"charged_weight",
]

REMITTANCE_COLUMNS = [
	"remittance_id",
	"remittance_date",
	"utr_number",
	"awb",
	"cod_amount",
	"freight_charge",
	"cod_fee",
	"gst",
	"net_paid",
	"charged_weight",
	"rto_charge",
	"weight_charge",
	"other_charge",
	"row_no",
]

# standard field -> header names (compared case-insensitively, ignoring spaces and punctuation)
ALIASES = {
	"Parcels": {
		"order_id": ["order id", "order no", "order number", "order name", "orderid"],
		"awb": [
			"awb",
			"awb number",
			"awb no",
			"waybill",
			"waybill number",
			"waybill no",
			"tracking number",
			"tracking id",
			"tracking no",
		],
		"courier": ["courier", "courier name", "courier partner", "carrier", "shipping partner"],
		"customer_name": ["customer name", "customer", "consignee name", "consignee"],
		"phone": ["phone", "mobile", "phone number", "customer phone"],
		"email": ["email", "customer email"],
		"address": ["address", "shipping address"],
		"city": ["city"],
		"pincode": ["pincode", "pin code", "zip", "postal code"],
		"cod_amount": ["cod amount", "cod", "cod value", "collectable amount", "amount to collect"],
		"order_total": ["order total", "order value", "total"],
		"payment_mode": ["payment mode", "payment method", "payment type"],
		"status": ["status", "shipment status", "delivery status"],
		"delivered_date": ["delivered date", "delivery date", "delivered on"],
		"order_date": ["order date", "created at", "ordered on"],
		"declared_weight": [
			"declared weight",
			"declared weight kg",
			"weight",
			"weight kg",
			"applied weight",
		],
		"charged_weight": [
			"charged weight",
			"charged weight kg",
			"charge weight",
			"billed weight",
			"chargeable weight",
			"chargeable weight kg",
		],
	},
	"Remittance": {
		"remittance_id": ["remittance id", "remittance no", "remittance number", "settlement id", "batch id"],
		"remittance_date": ["remittance date", "payment date", "settlement date", "paid date"],
		"utr_number": ["utr", "utr number", "utr no", "bank reference"],
		"awb": [
			"awb",
			"awb number",
			"awb no",
			"waybill",
			"waybill number",
			"waybill no",
			"tracking number",
			"tracking id",
			"tracking no",
		],
		"cod_amount": ["cod amount", "cod", "cod value", "collected amount"],
		"freight_charge": [
			"freight charges",
			"freight charge",
			"freight",
			"shipping charges",
			"shipping charge",
		],
		"cod_fee": ["cod charges", "cod charge", "cod fee", "cod handling fee"],
		"gst": ["gst", "gst amount", "tax", "igst"],
		"net_paid": [
			"net paid",
			"net amount",
			"net payable",
			"remitted amount",
			"amount paid",
			"net remittance",
		],
		"charged_weight": ["charged weight", "charge weight", "billed weight"],
		"rto_charge": ["rto charges", "rto charge"],
		"weight_charge": ["weight charges", "weight charge", "weight discrepancy"],
		"other_charge": ["other charges", "other charge", "other deductions"],
		"row_no": ["row no", "sr no", "s no", "serial no"],
	},
}

PARCEL_NUMERIC = ("cod_amount", "order_total")
PARCEL_WEIGHTS = ("declared_weight", "charged_weight")
REMITTANCE_NUMERIC = (
	"cod_amount",
	"freight_charge",
	"cod_fee",
	"gst",
	"net_paid",
	"charged_weight",
	"rto_charge",
	"weight_charge",
	"other_charge",
)
PARCEL_STATUSES = {
	"booked": "Booked",
	"intransit": "In Transit",
	"delivered": "Delivered",
	"rto": "RTO",
	"cancelled": "Cancelled",
	"canceled": "Cancelled",
}


def _key(text):
	return re.sub(r"[^a-z0-9]", "", str(text or "").lower())


def _check_roles():
	frappe.only_for(IMPORT_ROLES)


def _cell_to_str(value, fieldname=None):
	"""Convert an Excel/CSV cell value to a clean string (dates to ISO, whole floats without .0)."""
	if value is None:
		return ""
	if isinstance(value, datetime.datetime):
		return (
			value.strftime("%Y-%m-%d %H:%M:%S")
			if (value.hour or value.minute or value.second)
			else str(value.date())
		)
	if isinstance(value, datetime.date):
		return str(value)
	if isinstance(value, float):
		return str(int(value)) if value == int(value) else str(value)
	s = str(value).strip()
	if fieldname in ("awb", "phone", "pincode") and s.endswith(".0"):
		s = s[:-2]
	return s


def clean(value):
	return str(value).strip() if value not in (None, "") else ""


# ---------------------------------------------------------------------------
# reading and mapping
# ---------------------------------------------------------------------------


def _read_raw(file_url):
	"""Return all raw rows (list of lists) of a CSV/XLSX file."""
	path = frappe.get_doc("File", {"file_url": file_url}).get_full_path()

	if path.lower().endswith((".xlsx", ".xlsm")):
		from openpyxl import load_workbook

		wb = load_workbook(path, read_only=True, data_only=True)
		raw = list(wb.active.iter_rows(values_only=True))
		wb.close()
	elif path.lower().endswith(".xls"):
		import xlrd

		book = xlrd.open_workbook(path)
		sheet = book.sheet_by_index(0)
		raw = []
		for r in range(sheet.nrows):
			row = []
			for c in range(sheet.ncols):
				cell = sheet.cell(r, c)
				if cell.ctype == xlrd.XL_CELL_DATE:
					row.append(xlrd.xldate.xldate_as_datetime(cell.value, book.datemode))
				elif cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
					row.append(None)
				else:
					row.append(cell.value)
			raw.append(row)
	elif path.lower().endswith(".csv"):
		# path comes from the File record (get_full_path), not from user input
		with open(path, newline="", encoding="utf-8-sig") as f:  # nosemgrep
			raw = list(csv.reader(f))
	else:
		frappe.throw(_("Only CSV, XLSX or XLS files are supported"))
	return raw


def file_headers(file_url):
	"""Raw (lower-cased, trimmed) header cells of the file."""
	raw = _read_raw(file_url)
	return [str(h or "").strip().lower() for h in raw[0]] if raw else []


def _saved_mapping(courier, import_type, explicit=None):
	"""{file header (lower): standard field} from an explicit mapping or the Courier's saved JSON."""
	raw = explicit
	if not raw and courier and frappe.db.exists("Courier", courier):
		raw = frappe.db.get_value(
			"Courier", courier, "parcels_mapping" if import_type == "Parcels" else "remittance_mapping"
		)
	if isinstance(raw, str):
		try:
			raw = json.loads(raw) if raw.strip() else {}
		except ValueError:
			raw = {}
	if isinstance(raw, list):  # [{file_column, field}]
		raw = {r.get("file_column"): r.get("field") for r in raw if r.get("file_column") and r.get("field")}
	return {str(k).strip().lower(): v for k, v in (raw or {}).items() if v}


def build_mapping(headers, import_type, courier=None, explicit=None):
	"""Map each file header to a standard field.

	Saved/explicit mapping wins; headers it does not cover are auto-detected from common names.
	Returns [{"header", "field", "source"}] (source: saved / auto / none). A field is mapped once.
	"""
	saved = _saved_mapping(courier, import_type, explicit)
	by_alias = {}
	for field, names in ALIASES[import_type].items():
		for n in names:
			by_alias.setdefault(_key(n), field)
		by_alias.setdefault(_key(field), field)

	out, used = [], set()
	for h in headers:
		field, source = saved.get(h), "saved"
		if field and field in used:
			field = None
		if not field:
			field, source = by_alias.get(_key(h)), "auto"
			if field in used:
				field = None
		if field:
			used.add(field)
		else:
			source = "none"
		out.append({"header": h, "field": field, "source": source})
	return out


def _as_mapping(value):
	"""Accept a dict, a JSON string or None; return a dict (or list) for the mapping argument."""
	if value is None or value == "":
		return {}
	if isinstance(value, str):
		return frappe.parse_json(value) or {}
	return value


def _load(file_url, import_type, courier=None, mapping_template=None):
	"""Return (mapping, rows) where rows = [(excel_row_no, {field: str})] and unmapped columns are dropped."""
	raw = _read_raw(file_url)
	if not raw:
		return [], []
	headers = [str(h or "").strip().lower() for h in raw[0]]
	mapping = build_mapping(headers, import_type, courier, mapping_template)
	fields = [m["field"] for m in mapping]
	rows = []
	for i, values in enumerate(raw[1:], start=2):
		if not any(v not in (None, "") for v in values):
			continue
		row = {}
		for f, v in zip(fields, values, strict=False):
			if f:
				row[f] = _cell_to_str(v, fieldname=f)
		rows.append((i, row))
	return mapping, rows


def read_rows(file_url, mapping_template=None, courier=None, import_type="Parcels"):
	"""List of dicts keyed by standard field (kept for older callers)."""
	return [r for _i, r in _load(file_url, import_type, courier, mapping_template)[1]]


@frappe.whitelist()
def save_mapping(courier: str, import_type: str, mapping: dict | list | str):
	"""Save a {file header: standard field} mapping JSON on the Courier."""
	_check_roles()
	mapping = _as_mapping(mapping)
	if isinstance(mapping, list):
		mapping = {m.get("header") or m.get("file_column"): m.get("field") for m in mapping if m.get("field")}
	fieldname = "parcels_mapping" if import_type == "Parcels" else "remittance_mapping"
	frappe.db.set_value("Courier", courier, fieldname, json.dumps(mapping))
	return True


# ---------------------------------------------------------------------------
# row validation (pure: no writes)
# ---------------------------------------------------------------------------


def _num(row, field, errors, negative_ok=True):
	value = clean(row.get(field))
	if not value:
		return 0.0
	try:
		n = float(value.replace(",", ""))
	except ValueError:
		errors.append(_("{0} is not a number ({1})").format(field, value))
		return 0.0
	if n < 0 and not negative_ok:
		errors.append(_("{0} cannot be negative ({1})").format(field, value))
	return n


def _date(row, field, errors, as_datetime=False):
	value = clean(row.get(field))
	if not value:
		return None
	try:
		return str(get_datetime(value)) if as_datetime else str(getdate(value))
	except Exception:
		errors.append(_("{0} is not a valid date ({1})").format(field, value))
		return None


def _resolve_courier(name, cache):
	"""Courier record name for a file value (case-insensitive), or None."""
	key = (name or "").strip().lower()
	if not key:
		return None
	if not cache:
		cache.update({c.lower(): c for c in frappe.get_all("Courier", pluck="name")})
	return cache.get(key)


def check_parcel_row(row, default_courier, couriers):
	"""Return (values, errors) for one parcel row."""
	errors = []
	awb = normalize_awb(clean(row.get("awb")))
	order_id = clean(row.get("order_id")) or awb
	if not order_id:
		errors.append(_("Order ID or AWB is required"))
	v = {"order_id": order_id, "awb": awb or None}
	for f in PARCEL_NUMERIC:
		v[f] = _num(row, f, errors, negative_ok=False)
	for f in PARCEL_WEIGHTS:  # blank or non-numeric means "no weight", never an error
		try:
			w = float(clean(row.get(f)).replace(",", ""))
		except ValueError:
			w = None
		v[f] = w if w is not None and w > 0 else None
	v["delivered_date"] = _date(row, "delivered_date", errors)
	v["order_date"] = _date(row, "order_date", errors, as_datetime=True)

	courier_name = clean(row.get("courier"))
	if courier_name:
		v["courier"] = _resolve_courier(courier_name, couriers)
		if not v["courier"]:
			errors.append(_("Courier {0} not found").format(courier_name))
	else:
		v["courier"] = default_courier or None

	mode = clean(row.get("payment_mode")).lower()
	v["payment_mode"] = "Prepaid" if mode in ("prepaid", "paid") else "COD"
	status = clean(row.get("status"))
	if status:
		v["status"] = PARCEL_STATUSES.get(_key(status))
		if not v["status"]:
			errors.append(_("Status {0} is not valid").format(status))
	for f in ("customer_name", "phone", "email", "address", "city", "pincode"):
		v[f] = clean(row.get(f))
	return v, errors


def check_line_row(row, idx):
	"""Return (line values, errors) for one remittance row."""
	errors = []
	awb = normalize_awb(clean(row.get("awb")))
	if not awb:
		errors.append(_("AWB is required"))
	n = {f: _num(row, f, errors, negative_ok=(f != "net_paid")) for f in REMITTANCE_NUMERIC}
	if not clean(row.get("net_paid")):
		errors.append(_("net_paid is required"))
	return {
		"awb": awb,
		"cod_amount": n["cod_amount"],
		"freight": n["freight_charge"],
		"cod_fee": n["cod_fee"],
		"gst": n["gst"],
		"net_paid": n["net_paid"],
		"charged_weight": n["charged_weight"],
		"rto_charge": n["rto_charge"],
		"weight_charge": n["weight_charge"],
		"other_charge": n["other_charge"],
		"row_no": cint(clean(row.get("row_no"))) or idx,
	}, errors


def _batch_id(row):
	return clean(row.get("remittance_id")) or clean(row.get("utr_number"))


def _remittance_date(row, errors):
	d = _date(row, "remittance_date", errors)
	return d or nowdate()


# ---------------------------------------------------------------------------
# preview
# ---------------------------------------------------------------------------


def _line_detail(idx, v, errs, bid, seen_awbs, parcel_idx):
	"""One preview row with its validation label (no writes)."""
	deductions = sum(
		v[k] for k in ("freight", "cod_fee", "gst", "rto_charge", "weight_charge", "other_charge")
	)
	row = {
		"row": idx,
		"awb": v["awb"],
		"order": None,
		"cod_amount": v["cod_amount"],
		"deductions": round(deductions, 2),
		"net": v["net_paid"],
		"batch": bid,
		"detail": "",
	}
	if errs:
		return row | {"validation": "Invalid", "detail": "; ".join(errs)}
	matches = parcel_idx.get(v["awb"], [])
	if v["awb"] in seen_awbs:
		label, detail = "Duplicate AWB", _("AWB appears more than once in the file")
	elif not matches:
		label, detail = "AWB not found", _("No order has this AWB (becomes an Unknown AWB exception)")
	elif len(matches) > 1:
		label, detail = "Ambiguous AWB", _("{0} orders share this AWB").format(len(matches))
	elif matches[0].payment_mode == "Prepaid":
		label, detail = "Prepaid order", _("Prepaid orders are not reconciled")
	elif matches[0].settlement_batch:
		label, detail = "Already settled", _("Already in {0}").format(matches[0].settlement_batch)
	else:
		label, detail = "Matched", ""
	seen_awbs.add(v["awb"])
	if matches:
		row["order"] = matches[0].name
	return row | {"validation": label, "detail": detail}


@frappe.whitelist()
def preview_import(
	file_url: str, import_type: str, courier: str | None = None, mapping_template: str | dict | None = None
):
	"""Dry run: headers, mapping, first rows, valid / invalid counts with reasons. Saves nothing."""
	mapping_template = _as_mapping(mapping_template)
	_check_roles()
	if import_type not in ALIASES:
		frappe.throw(_("Import type must be Parcels or Remittance"))
	mapping, rows = _load(file_url, import_type, courier, mapping_template)
	couriers = {}
	invalid, valid, create, update = [], 0, 0, 0
	duplicates, seen_batches = [], {}
	existing = set(frappe.get_all("COD Parcel", pluck="name")) if import_type == "Parcels" else set()
	detail, seen_awbs, parcel_idx = [], set(), {}
	if import_type == "Remittance":
		for r in frappe.get_all(
			"COD Parcel",
			filters={"awb": ["is", "set"]},
			fields=["name", "awb", "settlement_batch", "payment_mode"],
			limit_page_length=0,
		):
			parcel_idx.setdefault(normalize_awb(r.awb), []).append(r)

	for idx, row in rows:
		if import_type == "Parcels":
			v, errs = check_parcel_row(row, courier, couriers)
			if not errs:
				if v["order_id"] in existing:
					update += 1
				else:
					create += 1
					existing.add(v["order_id"])
		else:
			_v, errs = check_line_row(row, idx)
			bid = _batch_id(row)
			if not bid:
				errs.append(_("Remittance ID (or UTR) is required"))
			elif not errs:
				_remittance_date(row, errs)
				seen_batches.setdefault(bid, 0)
				seen_batches[bid] += 1
			detail.append(_line_detail(idx, _v, errs, bid, seen_awbs, parcel_idx))
		if errs:
			invalid.append({"row": idx, "reason": "; ".join(errs)})
		else:
			valid += 1

	for bid in seen_batches:
		if frappe.db.exists("Courier Remittance", bid):
			duplicates.append(bid)

	out = {
		"import_type": import_type,
		"headers": [m["header"] for m in mapping],
		"mapping": mapping,
		"mapping_source": (
			"saved"
			if any(m["source"] == "saved" for m in mapping)
			else "auto"
			if any(m["source"] == "auto" for m in mapping)
			else "none"
		),
		"row_count": len(rows),
		"valid_count": valid,
		"invalid_count": len(invalid),
		"invalid_rows": invalid[:MAX_LISTED_ERRORS],
		"preview": [{"row": i, **r} for i, r in rows[:PREVIEW_ROWS]],
		"fields": [m["field"] for m in mapping if m["field"]],
		"background": len(rows) > BACKGROUND_ROWS,
	}
	if import_type == "Parcels":
		out |= {"will_create": create, "will_update": update}
	else:
		for d in detail:
			if d["validation"] == "Matched" and d["batch"] in duplicates:
				d["validation"], d["detail"] = (
					"Already imported",
					_("Remittance {0} exists").format(d["batch"]),
				)
		tiles = {"matched": 0, "awb_not_found": 0, "duplicate": 0, "invalid": 0, "other": 0}
		for d in detail:
			k = {
				"Matched": "matched",
				"AWB not found": "awb_not_found",
				"Invalid": "invalid",
				"Duplicate AWB": "duplicate",
				"Already imported": "duplicate",
				"Already settled": "duplicate",
			}.get(d["validation"], "other")
			tiles[k] += 1
		order = {
			"Invalid": 0,
			"AWB not found": 1,
			"Duplicate AWB": 2,
			"Ambiguous AWB": 2,
			"Already imported": 3,
			"Already settled": 3,
			"Prepaid order": 3,
			"Matched": 4,
		}
		detail.sort(key=lambda d: (order.get(d["validation"], 4), d["row"]))
		out |= {
			"batches": list(seen_batches),
			"duplicate_batches": duplicates,
			"tiles": tiles,
			"rows": detail[:MAX_DETAIL_ROWS],
			"rows_truncated": len(detail) > MAX_DETAIL_ROWS,
			"totals": {
				"cod": round(sum(d["cod_amount"] for d in detail), 2),
				"deductions": round(sum(d["deductions"] for d in detail), 2),
				"net": round(sum(d["net"] for d in detail), 2),
			},
		}
	if not rows:
		out["message"] = _("The file has no data rows")
	return out


# ---------------------------------------------------------------------------
# import
# ---------------------------------------------------------------------------


def _log(import_type, status, file_url, courier, payload, error_message=None):
	from codsync.codsync.synclog import log_sync

	log_sync(
		"Excel",
		"Excel Import",
		status,
		endpoint=file_url,
		error_message=error_message,
		reference_doctype="Courier" if courier else None,
		reference_name=courier or None,
		response_payload={"import_type": import_type, **payload},
	)


def _status(done, errors):
	return "Success" if (done or not errors) else "Failed"


UPDATABLE = (
	"customer_name",
	"phone",
	"email",
	"address",
	"city",
	"pincode",
	"order_date",
	"cod_amount",
	"order_total",
	"declared_weight",
	"charged_weight",
	"delivered_date",
)


def _upsert_parcel(v):
	"""Create one COD Parcel; for an existing order id only empty weights are backfilled."""
	name = v["order_id"]
	if frappe.db.exists("COD Parcel", name):
		doc = frappe.get_doc("COD Parcel", name)
		# an existing parcel is only backfilled: weights are filled when currently 0 or empty
		for f in PARCEL_WEIGHTS:
			if v.get(f) and not flt(doc.get(f)):
				doc.set(f, v[f])
		doc.save()
		return "updated"
	if v.get("awb"):
		other = frappe.db.get_value("COD Parcel", {"awb": v["awb"]})
		if other:
			frappe.throw(_("AWB {0} already belongs to order {1}").format(v["awb"], other))
	doc = {
		"doctype": "COD Parcel",
		"platform": "Excel",
		"order_id": name,
		"awb": v["awb"],
		"courier": v["courier"],
		"payment_mode": v["payment_mode"],
		"status": v.get("status") or "Booked",
	}
	for f in UPDATABLE:
		if v.get(f) not in (None, ""):
			doc[f] = v[f]
	frappe.get_doc(doc).insert()
	return "created"


@frappe.whitelist()
def import_parcels(file_url: str, courier: str | None = None, mapping_template: str | dict | None = None):
	"""Upsert COD Parcels by order id (platform Excel, payment mode COD unless the file says Prepaid)."""
	mapping_template = _as_mapping(mapping_template)
	_check_roles()
	try:
		_mapping, rows = _load(file_url, "Parcels", courier, mapping_template)
	except Exception as e:
		_log("Parcels", "Failed", file_url, courier, {}, error_message=str(e))
		raise

	couriers, errors, created, updated = {}, [], 0, 0
	for idx, row in rows:
		v, errs = check_parcel_row(row, courier, couriers)
		if errs:
			errors.append({"row": idx, "reason": "; ".join(errs)})
			continue
		frappe.db.savepoint("codsync_import")
		try:
			if _upsert_parcel(v) == "created":
				created += 1
			else:
				updated += 1
		except Exception as e:
			frappe.db.rollback(save_point="codsync_import")
			errors.append({"row": idx, "reason": frappe.utils.strip_html(str(e))})

	result = {
		"import_type": "Parcels",
		"rows": len(rows),
		"created": created,
		"updated": updated,
		"skipped": len(errors),
		"errors": errors[:MAX_LISTED_ERRORS],
	}
	_log(
		"Parcels",
		_status(created + updated, errors),
		file_url,
		courier,
		{k: v for k, v in result.items() if k != "import_type"},
		error_message=f"{len(errors)} row(s) skipped" if errors else None,
	)
	return result


@frappe.whitelist()
def import_remittance(file_url: str, courier: str, mapping_template: str | dict | None = None):
	"""Create + submit one Courier Remittance per remittance id, then run matching on each."""
	mapping_template = _as_mapping(mapping_template)
	_check_roles()
	from codsync.codsync.matching import run_matching

	if not courier:
		frappe.throw(_("Courier is required"))
	try:
		_mapping, rows = _load(file_url, "Remittance", courier, mapping_template)
	except Exception as e:
		_log("Remittance", "Failed", file_url, courier, {}, error_message=str(e))
		raise

	groups = {}
	errors = []
	for idx, row in rows:
		bid = _batch_id(row)
		if not bid:
			errors.append({"row": idx, "reason": _("Remittance ID (or UTR) is required")})
			continue
		groups.setdefault(bid, []).append((idx, row))

	file_name = frappe.db.get_value("File", {"file_url": file_url}, "file_name")
	result = {
		"import_type": "Remittance",
		"rows": len(rows),
		"remittances": [],
		"lines": 0,
		"duplicates": [],
		"matching": {},
		"errors": errors,
	}

	for bid, items in groups.items():
		if frappe.db.exists("Courier Remittance", bid):
			result["duplicates"].append(bid)
			continue
		lines = []
		for idx, row in items:
			line, errs = check_line_row(row, idx)
			if errs:
				errors.append({"row": idx, "reason": "; ".join(errs)})
			else:
				lines.append(line)
		if not lines:
			errors.append(
				{"row": items[0][0], "reason": _("{0}: no valid lines, nothing imported").format(bid)}
			)
			continue
		date_errs = []
		date = _remittance_date(items[0][1], date_errs)
		if date_errs:
			errors.append({"row": items[0][0], "reason": f"{bid}: " + "; ".join(date_errs)})
			continue

		frappe.db.savepoint("codsync_import")
		try:
			doc = frappe.get_doc(
				{
					"doctype": "Courier Remittance",
					"courier": courier,
					"remittance_id": bid,
					"remittance_date": date,
					"utr_number": clean(items[0][1].get("utr_number")),
					"source": "Excel",
					"file_name": file_name,
					"lines": lines,
				}
			)
			doc.insert()
			doc.submit()
			counts = run_matching(doc.name)
		except Exception as e:
			frappe.db.rollback(save_point="codsync_import")
			errors.append({"row": items[0][0], "reason": f"{bid}: " + frappe.utils.strip_html(str(e))})
			continue
		result["remittances"].append(doc.name)
		result["lines"] += len(lines)
		for k, n in (counts or {}).items():
			result["matching"][k] = result["matching"].get(k, 0) + n

	if result["duplicates"] and not result["remittances"]:
		result["message"] = _("Remittance {0} was already imported. Nothing was created.").format(
			", ".join(result["duplicates"])
		)
	result["skipped"] = len(errors)
	result["errors"] = errors[:MAX_LISTED_ERRORS]
	_log(
		"Remittance",
		_status(result["remittances"], errors),
		file_url,
		courier,
		{k: v for k, v in result.items() if k != "import_type"},
		error_message=result.get("message") or (f"{len(errors)} row(s) skipped" if errors else None),
	)
	return result


def _run_background(import_type, file_url, courier, mapping_template, user):
	"""Worker for large files: run the import, then tell the user's browser."""
	if import_type == "Parcels":
		result = import_parcels(file_url, courier, mapping_template)
	else:
		result = import_remittance(file_url, courier, mapping_template)
	# background job: nothing else commits the import
	frappe.db.commit()  # nosemgrep
	frappe.publish_realtime("codsync_import_done", result, user=user)


@frappe.whitelist(methods=["POST"])
def run_import(
	import_type: str,
	file_url: str,
	courier: str | None = None,
	mapping_template: str | dict | None = None,
	save_mapping_back: int = 0,
):
	"""Confirm step of the dialog. Optionally saves the detected mapping on the Courier first;
	large files are queued and finish with a codsync_import_done realtime event."""
	mapping_template = _as_mapping(mapping_template)
	_check_roles()
	if import_type not in ALIASES:
		frappe.throw(_("Import type must be Parcels or Remittance"))
	if import_type == "Remittance" and not courier:
		frappe.throw(_("Courier is required"))

	mapping, rows = _load(file_url, import_type, courier, mapping_template)
	if cint(save_mapping_back) and courier:
		save_mapping(courier, import_type, {m["header"]: m["field"] for m in mapping if m["field"]})

	if len(rows) > BACKGROUND_ROWS:
		frappe.enqueue(
			"codsync.codsync.importer._run_background",
			queue="long",
			timeout=3600,
			import_type=import_type,
			file_url=file_url,
			courier=courier,
			mapping_template=mapping_template,
			user=frappe.session.user,
		)
		return {"queued": True, "rows": len(rows)}
	if import_type == "Parcels":
		return import_parcels(file_url, courier, mapping_template)
	return import_remittance(file_url, courier, mapping_template)
